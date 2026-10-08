import logging
from logging.handlers import RotatingFileHandler

import pandas as pd
import pyodbc

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# إعدادات المشروع
# ============================================================

FILE_NAME = "customers_filtered.xlsx"
LOG_FILE = "cleaner.log"

MAX_LOG_SIZE = 1 * 1024 * 1024  # 1 MB
BACKUP_LOG_COUNT = 5


# ============================================================
# إعداد Logging
# ============================================================

def setup_logger():
    """إنشاء وإعداد نظام Logging مع تدوير الملفات."""

    logger = logging.getLogger("excel_data_cleaner")

    logger.setLevel(logging.INFO)

    # منع إضافة Handler أكثر من مرة
    if logger.handlers:
        return logger

    log_handler = RotatingFileHandler(
        LOG_FILE,
        maxBytes=MAX_LOG_SIZE,
        backupCount=BACKUP_LOG_COUNT,
        encoding="utf-8"
    )

    log_formatter = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(message)s"
    )

    log_handler.setFormatter(log_formatter)

    logger.addHandler(log_handler)

    return logger


logger = setup_logger()


# ============================================================
# الاتصال بقاعدة البيانات
# ============================================================

def connect_to_database():
    """الاتصال بـ SQL Server وإرجاع الاتصال."""

    try:
        connection = pyodbc.connect(
            "DRIVER={ODBC Driver 17 for SQL Server};"
            "SERVER=DESKTOP-C27FUFT;"
            "DATABASE=SalesDB;"
            "Trusted_Connection=yes;"
        )

        print("تم الاتصال بـ SQL Server بنجاح.")

        logger.info(
            "تم الاتصال بـ SQL Server بنجاح."
        )

        return connection

    except pyodbc.Error as error:

        print("❌ فشل الاتصال بـ SQL Server.")

        print(
            "تأكد من أن SQL Server يعمل "
            "وأن بيانات الاتصال صحيحة."
        )

        print(f"تفاصيل الخطأ: {error}")

        logger.error(
            "فشل الاتصال بـ SQL Server: %s",
            error,
            exc_info=True
        )

        return None


# ============================================================
# قراءة البيانات
# ============================================================

def read_data(connection):
    """قراءة البيانات من SQL Server وإرجاع DataFrame."""

    try:

        query = """
            SELECT Name, Email, Age
            FROM Customers
            WHERE Age >= 20
        """

        df = pd.read_sql(
            query,
            connection
        )

        print(
            "تمت قراءة البيانات من SQL Server بنجاح."
        )

        logger.info(
            "تمت قراءة البيانات بنجاح. عدد الصفوف: %d",
            len(df)
        )

        return df

    except Exception as error:

        print(
            "❌ حدث خطأ أثناء قراءة البيانات "
            "من SQL Server."
        )

        print(f"تفاصيل الخطأ: {error}")

        logger.error(
            "حدث خطأ أثناء قراءة البيانات: %s",
            error,
            exc_info=True
        )

        return None


# ============================================================
# تنظيف البيانات
# ============================================================

def clean_data(df, progress_callback=None):
    """تنظيف البيانات وإرجاع البيانات النظيفة وإحصائيات التنظيف."""

    def report_progress(value, message):
        if progress_callback:
            try:
                progress_callback(value, message)
            except Exception as callback_error:
                logger.warning(
                    "تعذر إرسال تحديث التقدم: %s",
                    callback_error
                )

    try:

        rows_before = len(df)

        report_progress(
            45,
            "جاري إزالة الصفوف المكررة..."
        )

        logger.info(
            "بدأ تنظيف البيانات. "
            "عدد الصفوف قبل التنظيف: %d",
            rows_before
        )

        # ----------------------------------------------------
        # إزالة الصفوف المكررة
        # ----------------------------------------------------

        duplicates_removed = df.duplicated().sum()

        df = df.drop_duplicates()

        report_progress(
            48,
            "جاري إزالة الصفوف ذات القيم الفارغة..."
        )

        logger.info(
            "تم حذف %d صف مكرر.",
            duplicates_removed
        )

        # ----------------------------------------------------
        # إزالة الصفوف ذات القيم الفارغة
        # ----------------------------------------------------

        missing_rows = (
            df[["Name", "Email", "Age"]]
            .isna()
            .any(axis=1)
            .sum()
        )

        df = df.dropna(
            subset=[
                "Name",
                "Email",
                "Age"
            ]
        )

        logger.info(
            "تم حذف %d صف يحتوي على قيم فارغة.",
            missing_rows
        )

        report_progress(
            51,
            "جاري تنظيف عناوين البريد الإلكتروني..."
        )

        # ----------------------------------------------------
        # تنظيف الإيميلات
        # ----------------------------------------------------

        df["Email"] = (
            df["Email"]
            .astype(str)
            .str.replace(
                r"\s+",
                "",
                regex=True
            )
            .str.lower()
        )

        logger.info(
            "تم تنظيف قيم البريد الإلكتروني."
        )

        report_progress(
            53,
            "جاري التحقق من صحة الإيميلات..."
        )

        # ----------------------------------------------------
        # التحقق من صحة الإيميلات
        # ----------------------------------------------------

        email_pattern = (
            r"^[^@\s]+@[^@\s]+\.[^@\s]+$"
        )

        valid_email = df["Email"].str.fullmatch(
            email_pattern,
            na=False
        )

        invalid_emails_removed = (
            ~valid_email
        ).sum()

        df = df[valid_email]

        rows_after = len(df)

        logger.info(
            "تم حذف %d صف يحتوي على "
            "بريد إلكتروني غير صحيح.",
            invalid_emails_removed
        )

        logger.info(
            "انتهى تنظيف البيانات. "
            "عدد الصفوف بعد التنظيف: %d",
            rows_after
        )

        report_progress(
            55,
            "اكتمل تنظيف البيانات."
        )

        print("تم تنظيف البيانات بنجاح.")

        cleaning_stats = {
            "rows_before": rows_before,
            "duplicates_removed": duplicates_removed,
            "missing_rows": missing_rows,
            "invalid_emails_removed": invalid_emails_removed,
            "rows_after": rows_after
        }

        return df, cleaning_stats

    except Exception as error:

        print(
            "❌ حدث خطأ أثناء تنظيف البيانات."
        )

        print(f"تفاصيل الخطأ: {error}")

        logger.error(
            "حدث خطأ أثناء تنظيف البيانات: %s",
            error,
            exc_info=True
        )

        return None, None


# ============================================================
# إنشاء تقرير التنظيف
# ============================================================

def create_cleaning_report(cleaning_stats):
    """إنشاء DataFrame يحتوي على تقرير التنظيف."""

    try:

        report = pd.DataFrame({
            "Metric": [
                "Rows before cleaning",
                "Duplicate rows removed",
                "Rows with missing values removed",
                "Invalid email rows removed",
                "Rows after cleaning"
            ],

            "Count": [
                cleaning_stats["rows_before"],
                cleaning_stats["duplicates_removed"],
                cleaning_stats["missing_rows"],
                cleaning_stats["invalid_emails_removed"],
                cleaning_stats["rows_after"]
            ]
        })

        logger.info(
            "تم إنشاء Cleaning Report بنجاح."
        )

        return report

    except Exception as error:

        print(
            "❌ حدث خطأ أثناء إنشاء تقرير التنظيف."
        )

        print(f"تفاصيل الخطأ: {error}")

        logger.error(
            "حدث خطأ أثناء إنشاء Cleaning Report: %s",
            error,
            exc_info=True
        )

        return None


# ============================================================
# حفظ البيانات في Excel
# ============================================================

def save_to_excel(df, report):
    """حفظ البيانات النظيفة والتقرير في ملف Excel."""

    try:

        with pd.ExcelWriter(
            FILE_NAME,
            engine="openpyxl"
        ) as writer:

            df.to_excel(
                writer,
                sheet_name="Cleaned Data",
                index=False
            )

            report.to_excel(
                writer,
                sheet_name="Cleaning Report",
                index=False
            )

        print(
            "تم حفظ ملف Excel بنجاح."
        )

        logger.info(
            "تم حفظ ملف Excel بنجاح: %s",
            FILE_NAME
        )

        return True

    except PermissionError:

        print(
            "❌ لا يمكن حفظ ملف Excel."
        )

        print(
            "تأكد أن ملف customers_filtered.xlsx "
            "مغلق في Excel ثم حاول مرة أخرى."
        )

        logger.error(
            "فشل حفظ Excel لأن الملف مفتوح أو محمي: %s",
            FILE_NAME,
            exc_info=True
        )

        return False

    except Exception as error:

        print(
            "❌ حدث خطأ أثناء حفظ ملف Excel."
        )

        print(f"تفاصيل الخطأ: {error}")

        logger.error(
            "حدث خطأ أثناء حفظ Excel: %s",
            error,
            exc_info=True
        )

        return False


# ============================================================
# تنسيق ملف Excel
# ============================================================

def format_excel():
    """تطبيق التنسيق الاحترافي على ملف Excel."""

    try:

        workbook = load_workbook(
            FILE_NAME
        )

        for sheet_name in [
            "Cleaned Data",
            "Cleaning Report"
        ]:

            worksheet = workbook[sheet_name]

            # ------------------------------------------------
            # تجميد الصف الأول
            # ------------------------------------------------

            worksheet.freeze_panes = "A2"

            # ------------------------------------------------
            # إضافة Filter
            # ------------------------------------------------

            worksheet.auto_filter.ref = (
                worksheet.dimensions
            )

            # ------------------------------------------------
            # تنسيق عناوين الأعمدة
            # ------------------------------------------------

            for cell in worksheet[1]:

                cell.font = Font(
                    bold=True,
                    color="FFFFFF"
                )

                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="1F4E78"
                )

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

            # ------------------------------------------------
            # ضبط عرض الأعمدة تلقائيًا
            # ------------------------------------------------

            for column in worksheet.columns:

                max_length = 0

                column_letter = (
                    get_column_letter(
                        column[0].column
                    )
                )

                for cell in column:

                    if cell.value is not None:

                        cell_length = len(
                            str(cell.value)
                        )

                        if cell_length > max_length:
                            max_length = cell_length

                worksheet.column_dimensions[
                    column_letter
                ].width = max_length + 3

        workbook.save(
            FILE_NAME
        )

        print(
            "تم تنسيق ملف Excel بنجاح."
        )

        logger.info(
            "تم تنسيق ملف Excel بنجاح. "
            "تم تطبيق Filter و Freeze Panes "
            "وضبط عرض الأعمدة."
        )

        return True

    except PermissionError:

        print(
            "❌ لا يمكن تعديل ملف Excel."
        )

        print(
            "تأكد أن ملف customers_filtered.xlsx "
            "مغلق في Excel ثم حاول مرة أخرى."
        )

        logger.error(
            "فشل تعديل Excel لأن الملف مفتوح أو محمي.",
            exc_info=True
        )

        return False

    except Exception as error:

        print(
            "❌ حدث خطأ أثناء تنسيق Excel."
        )

        print(f"تفاصيل الخطأ: {error}")

        logger.error(
            "حدث خطأ أثناء تنسيق Excel: %s",
            error,
            exc_info=True
        )

        return False


# ============================================================
# إغلاق الاتصال
# ============================================================

def close_connection(connection):
    """إغلاق اتصال SQL Server."""

    if connection is None:
        return

    try:

        connection.close()

        print(
            "تم إغلاق الاتصال بـ SQL Server."
        )

        logger.info(
            "تم إغلاق الاتصال بـ SQL Server بنجاح."
        )

    except Exception as error:

        print(
            "⚠️ تعذر إغلاق اتصال SQL Server "
            "بشكل طبيعي."
        )

        print(f"تفاصيل الخطأ: {error}")

        logger.warning(
            "تعذر إغلاق اتصال SQL Server "
            "بشكل طبيعي: %s",
            error,
            exc_info=True
        )


# ============================================================
# طباعة التقرير النهائي
# ============================================================

def print_final_report(cleaning_stats):
    """طباعة تقرير التنظيف النهائي في Terminal."""

    print()

    print(
        "========== تقرير تنظيف البيانات =========="
    )

    print(
        f"عدد الصفوف قبل التنظيف: "
        f"{cleaning_stats['rows_before']}"
    )

    print(
        f"عدد الصفوف المكررة المحذوفة: "
        f"{cleaning_stats['duplicates_removed']}"
    )

    print(
        f"عدد الصفوف ذات القيم الفارغة المحذوفة: "
        f"{cleaning_stats['missing_rows']}"
    )

    print(
        f"عدد الصفوف ذات الإيميلات غير الصحيحة "
        f"المحذوفة: "
        f"{cleaning_stats['invalid_emails_removed']}"
    )

    print(
        f"عدد الصفوف النهائي: "
        f"{cleaning_stats['rows_after']}"
    )

    print(
        f"تم حفظ البيانات النظيفة في: "
        f"{FILE_NAME}"
    )

    print(
        "=========================================="
    )

    print(
        "✅ تم تنفيذ البرنامج بنجاح."
    )


# ============================================================
# البرنامج الرئيسي
# ============================================================

def main(progress_callback=None):
    """تشغيل جميع مراحل المشروع بالترتيب."""

    def report_progress(value, message):
        if progress_callback:
            try:
                progress_callback(
                    value,
                    message
                )
            except Exception as callback_error:
                logger.warning(
                    "تعذر إرسال تحديث التقدم: %s",
                    callback_error
                )

    logger.info(
        "========== بدء تشغيل البرنامج =========="
    )

    connection = None

    try:

        # ----------------------------------------------------
        # بداية التشغيل
        # ----------------------------------------------------

        report_progress(
            0,
            "جاري بدء عملية التنظيف..."
        )

        # ----------------------------------------------------
        # 1. الاتصال بقاعدة البيانات
        # ----------------------------------------------------

        report_progress(
            10,
            "جاري الاتصال بقاعدة البيانات..."
        )

        connection = connect_to_database()

        if connection is None:
            return False

        report_progress(
            15,
            "تم الاتصال بقاعدة البيانات."
        )

        # ----------------------------------------------------
        # 2. قراءة البيانات
        # ----------------------------------------------------

        report_progress(
            25,
            "جاري قراءة البيانات من SQL Server..."
        )

        df = read_data(
            connection
        )

        if df is None:
            return False

        report_progress(
            35,
            "تمت قراءة البيانات بنجاح."
        )

        # ----------------------------------------------------
        # 3. تنظيف البيانات
        # ----------------------------------------------------

        report_progress(
            45,
            "جاري تنظيف البيانات..."
        )

        cleaned_df, cleaning_stats = clean_data(
            df,
            progress_callback=report_progress
        )

        if cleaned_df is None:
            return False

        # ----------------------------------------------------
        # 4. إنشاء تقرير التنظيف
        # ----------------------------------------------------

        report_progress(
            65,
            "جاري إنشاء تقرير التنظيف..."
        )

        report = create_cleaning_report(
            cleaning_stats
        )

        if report is None:
            return False

        report_progress(
            70,
            "تم إنشاء تقرير التنظيف."
        )

        # ----------------------------------------------------
        # 5. حفظ Excel
        # ----------------------------------------------------

        report_progress(
            80,
            "جاري حفظ ملف Excel..."
        )

        if not save_to_excel(
            cleaned_df,
            report
        ):
            return False

        report_progress(
            88,
            "تم حفظ ملف Excel."
        )

        # ----------------------------------------------------
        # 6. تنسيق Excel
        # ----------------------------------------------------

        report_progress(
            92,
            "جاري تنسيق ملف Excel..."
        )

        if not format_excel():
            return False

        report_progress(
            100,
            "اكتملت عملية التنظيف بنجاح."
        )

        # ----------------------------------------------------
        # 7. التقرير النهائي
        # ----------------------------------------------------

        print_final_report(
            cleaning_stats
        )

        logger.info(
            "اكتمل تشغيل البرنامج بنجاح. "
            "Rows before: %d | Rows after: %d",
            cleaning_stats["rows_before"],
            cleaning_stats["rows_after"]
        )

        return cleaning_stats

    except Exception as error:

        print(
            "❌ حدث خطأ غير متوقع أثناء تشغيل البرنامج."
        )

        print(
            f"تفاصيل الخطأ: {error}"
        )

        logger.critical(
            "حدث خطأ غير متوقع أثناء تشغيل البرنامج: %s",
            error,
            exc_info=True
        )

        return False

    finally:

        # إغلاق الاتصال مهما حدث
        close_connection(
            connection
        )

        logger.info(
            "========== انتهاء تشغيل البرنامج =========="
        )


# ============================================================
# نقطة بداية البرنامج
# ============================================================

if __name__ == "__main__":
    main()